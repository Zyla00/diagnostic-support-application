from django.contrib.auth.decorators import login_required
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from django.http import JsonResponse, HttpResponse
from django.db import transaction
from django.utils import timezone
from django.shortcuts import render, redirect
import pandas as pd
from django.contrib import messages
from decimal import Decimal
from .forms import UploadFileForm
from django.contrib.auth.mixins import LoginRequiredMixin
from django.views.generic import TemplateView, FormView, View
from django.shortcuts import get_object_or_404
from datetime import datetime
from .forms import (MoodScaleForm, MoodEmotionForm, MoodNoteForm, CoffeHabitForm, CigaretteHabitForm, SportsForm,
                    AlcoholHabitForm, SleepForm, CombinedDayForm, DateRangeForm)
from .models import MoodScale, MoodEmotion, MoodNote, Sleep, CoffeHabit, CigaretteHabit, Sports, AlcoholHabit, Day


from django.utils.decorators import method_decorator
from webApp.utils import role_required


@login_required
@role_required("patient")
def my_wellness_view(request):
    return render(request, "wellnessTracker/some_template.html")

@method_decorator([login_required, role_required("patient")], name='dispatch')
class HabbitView(LoginRequiredMixin, TemplateView):
    template_name = 'welnessTracker/homepage.html'
    login_url = '/login/'
    redirect_field_name = 'next'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        day = self.get_or_create_day(user)
        context.update({
            'mood_scale_form': MoodScaleForm(instance=day.mood_scale),
            'mood_emotion_form': MoodEmotionForm(instance=day.mood_emotion),
            'mood_note_form': MoodNoteForm(instance=day.mood_note),
            'sleep_form': SleepForm(instance=day.sleep),
            'coffee_form': CoffeHabitForm(instance=day.coffee_habit),
            'cigarette_form': CigaretteHabitForm(instance=day.cigarette_habit),
            'alcohol_form': AlcoholHabitForm(instance=day.alcohol_habit),
            'sports_form': SportsForm(instance=day.sports),
        })
        return context

    @staticmethod
    def get_or_create_day(user):
        today = timezone.now().date()
        with transaction.atomic():
            day, created = Day.objects.select_related(
                'mood_scale',
                'mood_emotion',
                'mood_note',
                'sleep',
                'coffee_habit',
                'cigarette_habit',
                'alcohol_habit',
                'sports'
            ).get_or_create(user=user, date=today)

            if created:
                day.mood_scale = MoodScale.objects.create()
                day.mood_emotion = MoodEmotion.objects.create()
                day.mood_note = MoodNote.objects.create()
                day.sleep = Sleep.objects.create()
                day.coffee_habit = CoffeHabit.objects.create()
                day.cigarette_habit = CigaretteHabit.objects.create()
                day.alcohol_habit = AlcoholHabit.objects.create()
                day.sports = Sports.objects.create()
                day.save()
            elif not all([
                day.mood_scale,
                day.mood_emotion,
                day.mood_note,
                day.sleep,
                day.coffee_habit,
                day.cigarette_habit,
                day.alcohol_habit,
                day.sports
            ]):
                if not day.mood_scale:
                    day.mood_scale = MoodScale.objects.create()
                if not day.mood_emotion:
                    day.mood_emotion = MoodEmotion.objects.create()
                if not day.mood_note:
                    day.mood_note = MoodNote.objects.create()
                if not day.sleep:
                    day.sleep = Sleep.objects.create()
                if not day.coffee_habit:
                    day.coffee_habit = CoffeHabit.objects.create()
                if not day.cigarette_habit:
                    day.cigarette_habit = CigaretteHabit.objects.create()
                if not day.alcohol_habit:
                    day.alcohol_habit = AlcoholHabit.objects.create()
                if not day.sports:
                    day.sports = Sports.objects.create()
                day.save()

        return day

    def post(self, request, *args, **kwargs):
        user = request.user
        day = self.get_or_create_day(user)
        form_id_list = request.POST.getlist('form_id')
        form_id = form_id_list[0] if form_id_list else None

        form = None
        if form_id == 'mood-scale-form':
            form = MoodScaleForm(request.POST, instance=day.mood_scale)
        elif form_id == 'mood-emotions-form':
            form = MoodEmotionForm(request.POST, instance=day.mood_emotion)
        elif form_id == 'mood-note-form':
            form = MoodNoteForm(request.POST, instance=day.mood_note)
        elif form_id == 'sleep-form':
            form = SleepForm(request.POST, instance=day.sleep)
        elif form_id == 'coffee-form':
            form = CoffeHabitForm(request.POST, instance=day.coffee_habit)
        elif form_id == 'cigarette-form':
            form = CigaretteHabitForm(request.POST, instance=day.cigarette_habit)
        elif form_id == 'alcohol-form':
            form = AlcoholHabitForm(request.POST, instance=day.alcohol_habit)
        elif form_id == 'sports-form':
            form = SportsForm(request.POST, instance=day.sports)

        if not form:
            return JsonResponse({'success': False, 'reason': 'No valid form found'})

        if not form.is_valid():
            return JsonResponse({'success': False, 'reason': 'Form validation failed'})

        if form_id == 'cigarette-form' and not self.validate_cigarettes(form.cleaned_data):
            return JsonResponse({'success': False, 'reason': 'Cigarette validation failed'})
        elif form_id == 'alcohol-form' and not self.validate_alcohol(form.cleaned_data):
            return JsonResponse({'success': False, 'reason': 'Alcohol validation failed'})
        elif form_id == 'sports-form' and not self.validate_exercise(form.cleaned_data):
            return JsonResponse({'success': False, 'reason': 'Exercise validation failed'})

        form.save()
        return JsonResponse({'success': True})

    @staticmethod
    def validate_cigarettes(data):
        return bool(data.get('cigarettes')) or not data.get('cigarette_type')

    @staticmethod
    def validate_alcohol(data):
        return bool(data.get('alcohol_amount')) or not data.get('alcohol_type')

    @staticmethod
    def validate_exercise(data):
        return bool(data.get('exercise_times')) or not data.get('exercise_type')

@method_decorator([login_required, role_required("patient")], name='dispatch')
class DayCreateEditView(FormView):
    template_name = 'welnessTracker/day_create_edit.html'
    form_class = CombinedDayForm

    def get_form_kwargs(self):
        kwargs = super(DayCreateEditView, self).get_form_kwargs()
        kwargs['user'] = self.request.user
        if self.kwargs.get('pk'):
            day = get_object_or_404(Day, pk=self.kwargs['pk'], user=self.request.user.id)
            kwargs['instance_id'] = day.id

            initial = {
                'date': day.date,
                'mood_scale': day.mood_scale.scale if day.mood_scale else None,
                'emotions': day.mood_emotion.emotions if day.mood_emotion else [],
                'note': day.mood_note.note if day.mood_note else '',
                'slept_scale': day.sleep.slept_scale if day.sleep else None,
                'coffee_amount': day.coffee_habit.coffee_amount if day.coffee_habit else None,
                'coffee_unit': day.coffee_habit.coffee_unit if day.coffee_habit else '',
                'cigarettes': day.cigarette_habit.cigarettes if day.cigarette_habit else None,
                'cigarette_type': day.cigarette_habit.cigarette_type if day.cigarette_habit else '',
                'alcohol_amount': day.alcohol_habit.alcohol_amount if day.alcohol_habit else None,
                'alcohol_unit': day.alcohol_habit.alcohol_unit if day.alcohol_habit else '',
                'alcohol_type': day.alcohol_habit.alcohol_type if day.alcohol_habit else [],
                'exercise_times': day.sports.exercise_times if day.sports else None,
                'exercise_unit': day.sports.exercise_unit if day.sports else '',
                'exercise_type': day.sports.exercise_type if day.sports else [],
            }
            kwargs['initial'] = initial
        return kwargs

    def form_valid(self, form):
        if self.kwargs.get('pk'):
            day = get_object_or_404(Day, pk=self.kwargs['pk'], user=self.request.user)
        else:
            day = Day(user=self.request.user)

        if form.cleaned_data['date']:
            day.date = form.cleaned_data['date']
            day.save()

        mood_scale_value = form.cleaned_data.get('mood_scale')
        if mood_scale_value is not None:
            day.mood_scale = MoodScale.objects.create(scale=mood_scale_value)

        emotions_value = form.cleaned_data.get('emotions')
        if emotions_value:
            day.mood_emotion = MoodEmotion.objects.create(emotions=emotions_value)

        note_value = form.cleaned_data.get('note')
        if note_value:
            day.mood_note = MoodNote.objects.create(note=note_value)

        slept_scale_value = form.cleaned_data.get('slept_scale')
        if slept_scale_value is not None:
            day.sleep = Sleep.objects.create(slept_scale=slept_scale_value)

        coffee_amount_value = form.cleaned_data.get('coffee_amount')
        coffee_unit_value = form.cleaned_data.get('coffee_unit')
        if coffee_amount_value is not None or coffee_unit_value:
            day.coffee_habit = CoffeHabit.objects.create(
                coffee_amount=coffee_amount_value,
                coffee_unit=coffee_unit_value
            )

        cigarettes_value = form.cleaned_data.get('cigarettes')
        cigarette_type_value = form.cleaned_data.get('cigarette_type')
        if cigarettes_value is not None or cigarette_type_value:
            day.cigarette_habit = CigaretteHabit.objects.create(
                cigarettes=cigarettes_value,
                cigarette_type=cigarette_type_value
            )

        alcohol_amount_value = form.cleaned_data.get('alcohol_amount')
        alcohol_unit_value = form.cleaned_data.get('alcohol_unit')
        alcohol_type_value = form.cleaned_data.get('alcohol_type')
        if alcohol_amount_value is not None or alcohol_unit_value or alcohol_type_value:
            day.alcohol_habit = AlcoholHabit.objects.create(
                alcohol_amount=alcohol_amount_value,
                alcohol_unit=alcohol_unit_value,
                alcohol_type=alcohol_type_value
            )

        exercise_times_value = form.cleaned_data.get('exercise_times')
        exercise_unit_value = form.cleaned_data.get('exercise_unit')
        exercise_type_value = form.cleaned_data.get('exercise_type')
        if exercise_times_value is not None or exercise_unit_value or exercise_type_value:
            day.sports = Sports.objects.create(
                exercise_times=exercise_times_value,
                exercise_unit=exercise_unit_value,
                exercise_type=exercise_type_value
            )

        day.save()

        day_data = {
            'id': day.id,
            'date': day.date.strftime('%d-%m-%Y'),
            'mood_scale': day.mood_scale.scale if day.mood_scale else None,
            'emotions': day.mood_emotion.emotions if day.mood_emotion else [],
            'note': day.mood_note.note if day.mood_note else '',
            'slept_scale': day.sleep.slept_scale if day.sleep else None,
            'coffee_amount': day.coffee_habit.coffee_amount if day.coffee_habit else None,
            'coffee_unit': day.coffee_habit.coffee_unit if day.coffee_habit else '',
            'cigarettes': day.cigarette_habit.cigarettes if day.cigarette_habit else None,
            'cigarette_type': day.cigarette_habit.cigarette_type if day.cigarette_habit else '',
            'alcohol_amount': day.alcohol_habit.alcohol_amount if day.alcohol_habit else None,
            'alcohol_unit': day.alcohol_habit.alcohol_unit if day.alcohol_habit else '',
            'alcohol_type': day.alcohol_habit.alcohol_type if day.alcohol_habit else [],
            'exercise_times': day.sports.exercise_times if day.sports else None,
            'exercise_unit': day.sports.exercise_unit if day.sports else '',
            'exercise_type': day.sports.exercise_type if day.sports else [],
        }

        return JsonResponse({'success': True, 'day': day_data})

    def form_invalid(self, form):
        errors = {
            str(form.fields[field_name].label if field_name in form.fields and form.fields[field_name].label else field_name):
            [str(message) for message in error_messages]
            for field_name, error_messages in form.errors.items()
        }

        return JsonResponse({'success': False, 'errors': errors})

@method_decorator([login_required, role_required("patient")], name='dispatch')
class DayDeleteView(LoginRequiredMixin, View):
    login_url = '/login/'
    redirect_field_name = 'next'

    def post(self, request, *args, **kwargs):
        day_id = kwargs.get('pk')
        user = request.user
        day = get_object_or_404(Day, pk=day_id, user=user)

        if day:
            day.delete()
            return JsonResponse({'success': True})
        else:
            return JsonResponse({'success': False, 'error': 'Delete unsuccessful. Please try again later'})
@method_decorator([login_required, role_required("patient")], name='dispatch')
class FetchPreviousDayView(LoginRequiredMixin, View):
    login_url = '/login/'
    redirect_field_name = 'next'

    def get(self, request, *args, **kwargs):
        last_date_str = request.GET.get('last_date')

        if last_date_str:
            try:
                last_date = datetime.strptime(last_date_str, "%B %d, %Y").date()
            except ValueError:
                return JsonResponse({'error': 'Invalid date format'}, status=400)

            previous_day = Day.objects.filter(user=request.user, date__lt=last_date).order_by('-date').first()
            if previous_day:
                day_data = {
                    'id': previous_day.id,
                    'date': previous_day.date.strftime('%d-%m-%Y'),
                    'mood_scale': previous_day.mood_scale.scale if previous_day.mood_scale else None,
                    'emotions': previous_day.mood_emotion.emotions if previous_day.mood_emotion else [],
                    'note': previous_day.mood_note.note if previous_day.mood_note else '',
                    'slept_scale': previous_day.sleep.slept_scale if previous_day.sleep else None,
                    'coffee_amount': previous_day.coffee_habit.coffee_amount if previous_day.coffee_habit else None,
                    'coffee_unit': previous_day.coffee_habit.coffee_unit if previous_day.coffee_habit else '',
                    'cigarettes': previous_day.cigarette_habit.cigarettes if previous_day.cigarette_habit else None,
                    'cigarette_type': previous_day.cigarette_habit.cigarette_type if previous_day.cigarette_habit else '',
                    'alcohol_amount': previous_day.alcohol_habit.alcohol_amount if previous_day.alcohol_habit else None,
                    'alcohol_unit': previous_day.alcohol_habit.alcohol_unit if previous_day.alcohol_habit else '',
                    'alcohol_type': previous_day.alcohol_habit.alcohol_type if previous_day.alcohol_habit else [],
                    'exercise_times': previous_day.sports.exercise_times if previous_day.sports else None,
                    'exercise_unit': previous_day.sports.exercise_unit if previous_day.sports else '',
                    'exercise_type': previous_day.sports.exercise_type if previous_day.sports else [],
                }
                return JsonResponse({'day': day_data})
            else:
                return JsonResponse({'message': 'No more days available'}, status=404)
        else:
            return JsonResponse({'error': 'No date provided'}, status=400)

@method_decorator([login_required, role_required("patient")], name='dispatch')
class ExportDaysExcelView(LoginRequiredMixin, View):
    login_url = '/login/'
    redirect_field_name = 'next'

    def get(self, request, *args, **kwargs):
        date_from = request.GET.get('date_from')
        date_to = request.GET.get('date_to')

        if not date_from or not date_to:
            return HttpResponse("Invalid date range", status=400)

        try:
            date_from = datetime.strptime(date_from, '%Y-%m-%d').date()
            date_to = datetime.strptime(date_to, '%Y-%m-%d').date()
        except ValueError:
            return HttpResponse("Invalid date format", status=400)

        days = Day.objects.filter(user=request.user, date__range=(date_from, date_to)).order_by('date')

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Days Data"

        headers = ['Date', 'Mood Scale', 'Emotions', 'Note', 'Slept Scale', 'Coffee Amount', 'Coffee Unit',
                   'Cigarettes', 'Cigarette Type', 'Alcohol Amount', 'Alcohol Unit', 'Alcohol Type',
                   'Exercise Times', 'Exercise Unit', 'Exercise Type']
        worksheet.append(headers)

        column_widths = [15, 10, 50, 50, 10, 15, 10, 10, 15, 15, 10, 50, 15, 10, 50]
        for i, column_width in enumerate(column_widths, 1):
            worksheet.column_dimensions[get_column_letter(i)].width = column_width

        for day in days:
            worksheet.append([
                day.date.strftime('%Y-%m-%d'),
                day.mood_scale.scale if day.mood_scale else '',
                ", ".join(day.mood_emotion.emotions) if day.mood_emotion else '',
                day.mood_note.note if day.mood_note else '',
                day.sleep.slept_scale if day.sleep else '',
                day.coffee_habit.coffee_amount if day.coffee_habit else '',
                day.coffee_habit.coffee_unit if day.coffee_habit else '',
                day.cigarette_habit.cigarettes if day.cigarette_habit else '',
                day.cigarette_habit.cigarette_type if day.cigarette_habit else '',
                day.alcohol_habit.alcohol_amount if day.alcohol_habit else '',
                day.alcohol_habit.alcohol_unit if day.alcohol_habit else '',
                ", ".join(day.alcohol_habit.alcohol_type) if day.alcohol_habit else '',
                day.sports.exercise_times if day.sports else '',
                day.sports.exercise_unit if day.sports else '',
                ", ".join(day.sports.exercise_type) if day.sports else ''
            ])

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="days_{date_from}_{date_to}.xlsx"'

        workbook.save(response)
        return response

@method_decorator([login_required, role_required("patient")], name='dispatch')
class ProcessDaysExportView(LoginRequiredMixin, FormView):
    template_name = 'welnessTracker/day_download.html'
    form_class = DateRangeForm
    login_url = '/login/'
    redirect_field_name = 'next'

    def form_valid(self, form):
        date_from = form.cleaned_data['date_from']
        date_to = form.cleaned_data['date_to']
        return JsonResponse({'success': True, 'date_from': date_from.strftime('%Y-%m-%d'),
                             'date_to': date_to.strftime('%Y-%m-%d')})

    def form_invalid(self, form):
        errors = {
            str(form.fields[field_name].label if field_name in form.fields and form.fields[
                field_name].label else field_name):
                [str(message) for message in error_messages]
            for field_name, error_messages in form.errors.items()
        }

        return JsonResponse({'success': False, 'errors': errors})

@login_required
@role_required("patient")
def upload_file(request):
    if request.method == 'POST':
        form = UploadFileForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            df = pd.read_excel(file)

            for index, row in df.iterrows():
                mood_scale = MoodScale.objects.create(scale=row['mood_scale'])
                mood_emotion = MoodEmotion.objects.create(emotions=row['mood_emotion'].split(','))
                mood_note = MoodNote.objects.create(note=row['mood_note'])
                sleep = Sleep.objects.create(slept_scale=Decimal(row['sleep']))
                coffee_habit = CoffeHabit.objects.create(coffee_amount=row['coffee_amount'],
                                                         coffee_unit=row['coffee_unit'])
                cigarette_habit = CigaretteHabit.objects.create(cigarettes=row['cigarettes'],
                                                                cigarette_type=row['cigarette_type'])
                alcohol_habit = AlcoholHabit.objects.create(alcohol_amount=row['alcohol_amount'],
                                                            alcohol_unit=row['alcohol_unit'],
                                                            alcohol_type=row['alcohol_type'].split(','))
                sports = Sports.objects.create(exercise_times=row['exercise_times'], exercise_unit=row['exercise_unit'],
                                               exercise_type=row['exercise_type'].split(','))

                Day.objects.create(
                    user=request.user,
                    mood_scale=mood_scale,
                    mood_emotion=mood_emotion,
                    mood_note=mood_note,
                    sleep=sleep,
                    coffee_habit=coffee_habit,
                    cigarette_habit=cigarette_habit,
                    alcohol_habit=alcohol_habit,
                    sports=sports,
                    date=row['date']
                )
            messages.success(request, 'Data uploaded successfully')
            return redirect('view_name')  # Redirect to the view displaying the data
    else:
        form = UploadFileForm()
    return render(request, 'upload.html', {'form': form})


